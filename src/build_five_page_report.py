"""Revise the original narrative report using notebook-verified results; five-page layout."""
from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'tmp/sample_selection'
data=json.loads((TMP/'report_content.json').read_text(encoding='utf-8'))
doc=Document(ROOT/'reports/peinado_1012953314_assignment1_revised.docx')
ps=list(doc.paragraphs)
for i,text in data['paragraphs'].items():ps[int(i)].text=text
# Preserve the original narrative arc while replacing affected results and interpretations.
updates={
4:'How does developer pay vary with work arrangement and education? In the selected sample, remote respondents earn more at both the mean and median, while US/Canada degree groups differ in average pay. The question is how robust these patterns are to salary extremes and statistical uncertainty.',
5:'Within this population, satisfaction analysis omits 3,428 missing ratings and one value of 60, using a working scale of 0–10. Each row has one supplied job category. All eligible salary extremes remain in the primary analyses.',
7:'Figures 1–2 show the lowest two and highest three means; the notebook covers all 31 retained titles. Senior executives lead satisfaction (8.12/10, n=168). Financial analysts/engineers average $501,248 (n=46), against a median of $87,248.50: an immediate warning that a high mean can reflect a small upper tail.',
9:'Figure 3 retains salaries from $102 to $33.55 million on a log axis. Three recorded ages below 15 remain flagged in the notebook; their validity and the age-cleaning history are uncertain. These diagnostics cause no further salary exclusions.',
12:'The first comparison asks whether work arrangement is associated with pay. Exact labels Remote and Hybrid retain 14,664 respondents; 5,648 other eligible respondents are outside Q2. Positive Remote − Hybrid differences indicate higher remote pay.',
18:'Is the mean gap consistent with equal population means? I test H₀: μR − μH = 0 against H₁: μR − μH ≠ 0 at α=0.05. Sample variances, using n−1, are 197,358,671,571.47 and 120,149,529,197.14 USD². The pooled calculation is:',
26:'Bootstrapping offers a second view of uncertainty. Within each group, I resample with replacement, retaining 6,888 Remote and 7,776 Hybrid salaries per draw. With 10,000 replications, seed 16242026 and separate named NumPy streams, the mean SEs are $5,362.02 and $3,890.15; the difference SE is $6,605.63. These describe uncertainty in means, not dispersion between individual salaries.',
36:'Both 95% intervals exclude zero, supporting a positive mean gap. The pooled interval is $17,243.91 ± $12,827.04, using t(0.975, 14,662). The bootstrap interval uses the 2.5th and 97.5th percentiles of unshifted differences; it is $223.81 wider and asymmetric. Under their assumptions, these procedures target 95% coverage across repeated samples, not a 95% probability for the fixed population gap. Percentile intervals and t-tests need not always agree.',
38:'Does the advantage persist for a typical salary? Remote and Hybrid medians are $90,491 and $77,730, a gap of $12,761. Mood’s test counts values above versus at/below the pooled median of $81,870; ties (11 Remote, 5 Hybrid) go below. With Yates correction, χ²(1)=109.1245, p=1.524e-25, rejecting equal medians at 0.05.',
43:'Education provides a second perspective. The selected sample contains 5,413 US/Canada respondents; exact Bachelor’s, Master’s and Professional degree labels retain 4,316. B, M and P denote these groups below. Professional is used as supplied, without inferring its qualifications. All salaries in these eligible groups remain.',
45:'The degree means rise from B to P, but are they driven by the upper tail? A separate 1.5×IQR sensitivity check flags 151/54/21 salaries, leaving 2,803/1,013/274. Means fall to $142,251.76/$160,662.59/$160,710.76: the P–M gap narrows to just $48.17. Restricted ANOVA remains significant (p=9.382e-16). Thus the broad pattern survives, while the apparent separation of the two higher degree means is sensitive to extreme salaries.',
47:'To test the overall pattern, H₀: μB=μM=μP is compared with the alternative that at least one mean differs. ANOVA gives F(2, 4,313)=18.7362, p=7.909e-09. The exact reference assumes independence, normal errors and equal variances. A variance ratio of 2.96 and Levene p=1.626e-05 warrant caution; Welch ANOVA also rejects, F(2, 682.27)=12.9764, p=2.944e-06. Neither omnibus result identifies significant pairs.',
49:'The bootstrap reveals why the degree means have unequal precision. Across 10,000 resamples of sizes 2,954/1,067/295 (B/M/P), mean SEs are $2,100.56/$5,000.74/$11,546.91. The smaller Professional group is much less precise. Separate streams use seed 16242026; shared group draws preserve dependence between contrasts. Observed B−M, B−P and M−P gaps are −$22,122.51, −$38,585.18 and −$16,462.67.',
53:'ANOVA on the three arrays of 10,000 unshifted bootstrap means gives F(2, 29,997)=69,977.6132; its p-value underflows numerically. Although it rejects, increasing B adds simulations rather than respondents. As in Q2, this calculation supplies no independent confirmation and identifies no significant pair.',
54:'For a separate test under equal means, I shift each degree group to the combined mean and resample at its original size, preserving empirical variances and shapes. None of 10,000 null F* values reaches the observed statistic: plus-one p=1/10,001=0.000100. This approximate calibration also rejects, but the simulation floor is not a precise tail estimate. No formal post-hoc tests were performed.',
56:'The evidence tells a consistent but qualified story: remote respondents have higher mean and median pay, and US/Canada degree groups differ overall. Median and sensitivity results show that the remote advantage is not solely an extreme-salary phenomenon; degree-tail sensitivity and unequal bootstrap precision qualify the education ranking. These associations cannot isolate effects of work mode or education from country, occupation and experience. The selected sample is not a verified employed-only population. Self-report, survey selection and uncertain independence limit generalization; bootstrapping cannot repair them. Eight identical rows remain because respondent identities are unverified, and no overall multiplicity correction was applied.',
58:'OpenAI Codex assisted with analysis, drafting and revision. Prompts included “Create the written report”, “reduce it to 5 page”, and the current request to apply population criteria while preserving the narrative. Notebook checks cover manual/library agreement, selection and bootstrap calculations; all 28 code cells completed without errors. Independent student review and Colab execution remain unverified.',
59:'Sources: supplied CSV and executed notebook; SciPy test documentation and Shalizi, The Bootstrap §5.2.4 (links in notebook).'
}
for i,text in updates.items():ps[i].text=text
# Remove the now-outdated dashboard link rather than implying it reflects this sample.
ps[2]._element.getparent().remove(ps[2]._element)
heading=ps[5].insert_paragraph_before('Study Population and Sample Selection Criteria','Heading 2')
selection='To focus the study on workforce participation, I exclude DevType = “Retired” and recorded Age >80 before all analyses. These researcher-defined scope decisions reflect engineering judgment, not data cleaning or evidence that older people are outliers; age 80 remains eligible. Of 20,407 original rows, 2 retired and 93 older respondents are excluded, with no overlap, leaving 20,312. Both retired records also say Employed, so the explicit DevType classification governs. Retirement exclusion does not establish active employment: 321 students and 58 respondents marked Not employed remain. Supplementary paid-work responses support self-reported participation, not independent verification. The 10 missing ages remain outside the age plot, so the ceiling cannot be verified for every respondent.'
ps[5].insert_paragraph_before(selection,'Normal')
for table,rows in zip(doc.tables,data['tables']):
 for row,values in zip(table.rows[1:],rows):
  for cell,value in zip(row.cells,values):cell.text=value
for i,name,height in [(6,'eda',1.7),(8,'age',1.25),(27,'q2',1.9),(50,'q3',1.9)]:
 ps[i].clear();ps[i].add_run().add_picture(str(TMP/f'{name}.png'),width=Inches(6.5),height=Inches(height))
# Match the baseline's compact report typography with explicit values for reproducible pagination.
for name,size in [('Normal',9.5),('Caption',8.5),('Heading 1',12),('Heading 2',10.5),('Title',16)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.0
for p in doc.paragraphs:
 p.paragraph_format.space_before=Pt(0)
 p.paragraph_format.space_after=Pt(5)
 p.paragraph_format.line_spacing=1.0
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for run in p.runs:run.font.size=Pt(9)
path=ROOT/'reports/peinado_1012953314_assignment1_final_5pages.docx'
doc.save(path)
print(path)
